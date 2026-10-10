import csv
import io
import os


COLUMN_ALIASES = {
    "query_text": ["query_text", "query", "search_term", "search term", "search query", "keyword", "term"],
    "match_type": ["match_type", "match type", "matchtype"],
    "campaign": ["campaign", "campaign_name", "campaign name"],
    "ad_group": ["ad_group", "ad group", "adgroup", "ad_group_name"],
    "impressions": ["impressions", "imps", "views"],
    "clicks": ["clicks", "click_count"],
    "cost": ["cost", "spend", "ad_spend", "cost ($)", "cost_usd"],
    "conversions": ["conversions", "conv", "conversion_count", "orders"],
    "conversion_value": ["conversion_value", "revenue", "conv_value", "total_conv_value", "value ($)"]
}

REQUIRED_CANONICAL = ["query_text"]
OPTIONAL_METRICS = ["impressions", "clicks", "cost", "conversions", "conversion_value", "match_type", "campaign", "ad_group"]

class CSVProcessor:
    def __init__(self):
        pass

    @staticmethod
    def detect_and_map_columns(df_columns, custom_mapping=None):
        mapping = {}
        unmapped_raw = list(df_columns)
        
        if custom_mapping:
            for canon, raw in custom_mapping.items():
                if raw in df_columns and raw is not None:
                    mapping[canon] = raw
                    if raw in unmapped_raw:
                        unmapped_raw.remove(raw)

        for canon, aliases in COLUMN_ALIASES.items():
            if canon not in mapping:
                for col in df_columns:
                    if col.lower().strip() in [a.lower() for a in aliases]:
                        mapping[canon] = col
                        if col in unmapped_raw:
                            unmapped_raw.remove(col)
                        break

        missing_required = [c for c in REQUIRED_CANONICAL if c not in mapping]
        available_optional = [c for c in OPTIONAL_METRICS if c in mapping]
        missing_optional = [c for c in OPTIONAL_METRICS if c not in mapping]

        return {
            "canonical_mapping": mapping,
            "missing_required": missing_required,
            "available_optional": available_optional,
            "missing_optional": missing_optional,
            "unmapped_raw_columns": unmapped_raw
        }

    @classmethod
    def load_and_standardize_csv(cls, file_or_path, custom_mapping=None):
        if isinstance(file_or_path, str):
            with open(file_or_path, mode='r', encoding='utf-8-sig', errors='replace') as f:
                lines = f.readlines()
        else:
            content = file_or_path.read()
            if isinstance(content, bytes):
                content = content.decode('utf-8-sig', errors='replace')
            lines = content.splitlines()

        if not lines:
            raise ValueError("Uploaded CSV file is empty.")

        reader = csv.reader(lines)
        try:
            raw_headers = next(reader)
        except StopIteration:
            raise ValueError("Uploaded CSV file is empty.")

        col_analysis = cls.detect_and_map_columns(raw_headers, custom_mapping)
        if col_analysis["missing_required"]:
            raise ValueError(f"Missing required columns: {col_analysis['missing_required']}. Please map them manually.")

        mapping = col_analysis["canonical_mapping"]
        
        records = []
        valid_count = 0

        for idx, row in enumerate(reader):
            if not row:
                continue
            row_dict = {raw_headers[i]: row[i] if i < len(row) else "" for i in range(len(raw_headers))}
            
            std_rec = {
                "source_row_id": f"ROW_{idx+1:06d}",
                "query_text": row_dict.get(mapping.get("query_text", ""), "").strip(),
                "match_type": row_dict.get(mapping.get("match_type", ""), "Unknown").strip() or "Unknown",
                "campaign": row_dict.get(mapping.get("campaign", ""), "Default Campaign").strip() or "Default Campaign",
                "ad_group": row_dict.get(mapping.get("ad_group", ""), "Default Ad Group").strip() or "Default Ad Group"
            }

            if std_rec["query_text"]:
                valid_count += 1

            # Numeric conversion
            for num_col in ["impressions", "clicks", "cost", "conversions", "conversion_value"]:
                val_str = row_dict.get(mapping.get(num_col, ""), "").strip()
                if val_str != "":
                    try:
                        std_rec[num_col] = float(val_str)
                    except ValueError:
                        std_rec[num_col] = None
                else:
                    std_rec[num_col] = None

            # Preserve original extra columns
            for col in raw_headers:
                if col not in std_rec and col != "_synthetic_ground_truth":
                    std_rec[f"orig_{col}"] = row_dict.get(col, "")

            if "_synthetic_ground_truth" in row_dict:
                std_rec["_synthetic_ground_truth"] = row_dict["_synthetic_ground_truth"]

            records.append(std_rec)

        summary = {
            "total_rows": len(records),
            "valid_rows": valid_count,
            "invalid_rows": len(records) - valid_count,
            "mapped_columns": mapping,
            "missing_optional": col_analysis["missing_optional"],
            "available_optional": col_analysis["available_optional"]
        }

        return records, summary

    @classmethod
    def load_excel_headers_and_rows(cls, file_bytes, filename, sheet_name=None):
        """
        Loads headers and raw rows from an Excel file (.xls or .xlsx).
        Returns (sheet_names, selected_sheet_name, raw_headers, list_of_row_dicts).
        """
        ext = os.path.splitext(filename)[1].lower()
        sheet_names = []
        raw_headers = []
        rows = []

        if ext == ".xls":
            import xlrd
            workbook = xlrd.open_workbook(file_contents=file_bytes)
            sheet_names = workbook.sheet_names()
            sel_sheet = sheet_name if (sheet_name and sheet_name in sheet_names) else sheet_names[0]
            sheet = workbook.sheet_by_name(sel_sheet)

            if sheet.nrows > 0:
                raw_headers = [str(sheet.cell_value(0, col)).strip() for col in range(sheet.ncols)]
                for r in range(1, sheet.nrows):
                    row_vals = [sheet.cell_value(r, col) for col in range(sheet.ncols)]
                    row_dict = {raw_headers[c]: str(row_vals[c]).strip() for c in range(len(raw_headers))}
                    rows.append(row_dict)

        elif ext in [".xlsx", ".xlsm"]:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            sheet_names = wb.sheetnames
            sel_sheet = sheet_name if (sheet_name and sheet_name in sheet_names) else sheet_names[0]
            sheet = wb[sel_sheet]

            all_rows = list(sheet.iter_rows(values_only=True))
            if all_rows:
                raw_headers = [str(cell).strip() if cell is not None else "" for cell in all_rows[0]]
                for r in all_rows[1:]:
                    if not any(r):
                        continue
                    row_dict = {raw_headers[c]: str(r[c]).strip() if c < len(r) and r[c] is not None else "" for c in range(len(raw_headers))}
                    rows.append(row_dict)

        else:
            raise ValueError(f"Unsupported file format '{ext}'. Expected .csv, .xls, or .xlsx.")

        return sheet_names, sel_sheet, raw_headers, rows

    @classmethod
    def load_and_standardize_file(cls, file_bytes, filename, sheet_name=None, custom_mapping=None):
        """
        Unified loader for .csv, .xls, and .xlsx.
        """
        ext = os.path.splitext(filename)[1].lower()
        if ext in [".xls", ".xlsx", ".xlsm"]:
            sheet_names, sel_sheet, raw_headers, raw_rows = cls.load_excel_headers_and_rows(file_bytes, filename, sheet_name)
            
            if not raw_headers:
                raise ValueError(f"Worksheet '{sel_sheet}' in '{filename}' is empty or has no header row.")

            col_analysis = cls.detect_and_map_columns(raw_headers, custom_mapping)
            if col_analysis["missing_required"]:
                raise ValueError(f"Missing required columns: {col_analysis['missing_required']}. Please map them.")

            mapping = col_analysis["canonical_mapping"]
            records = []
            valid_count = 0

            for idx, row_dict in enumerate(raw_rows):
                std_rec = {
                    "source_row_id": f"ROW_{idx+1:06d}",
                    "query_text": row_dict.get(mapping.get("query_text", ""), "").strip(),
                    "match_type": row_dict.get(mapping.get("match_type", ""), "Unknown").strip() or "Unknown",
                    "campaign": row_dict.get(mapping.get("campaign", ""), "Default Campaign").strip() or "Default Campaign",
                    "ad_group": row_dict.get(mapping.get("ad_group", ""), "Default Ad Group").strip() or "Default Ad Group"
                }

                if std_rec["query_text"]:
                    valid_count += 1

                for num_col in ["impressions", "clicks", "cost", "conversions", "conversion_value"]:
                    val_str = row_dict.get(mapping.get(num_col, ""), "").strip()
                    if val_str != "" and val_str != "None":
                        try:
                            std_rec[num_col] = float(val_str)
                        except ValueError:
                            std_rec[num_col] = None
                    else:
                        std_rec[num_col] = None

                for col in raw_headers:
                    if col not in std_rec:
                        std_rec[f"orig_{col}"] = row_dict.get(col, "")

                records.append(std_rec)

            summary = {
                "filename": filename,
                "sheet_names": sheet_names,
                "selected_sheet": sel_sheet,
                "total_rows": len(records),
                "valid_rows": valid_count,
                "invalid_rows": len(records) - valid_count,
                "mapped_columns": mapping,
                "missing_optional": col_analysis["missing_optional"],
                "available_optional": col_analysis["available_optional"]
            }
            return records, summary

        else:
            # Fallback to standard CSV handling
            bio = io.BytesIO(file_bytes) if isinstance(file_bytes, bytes) else file_bytes
            return cls.load_and_standardize_csv(bio, custom_mapping)


