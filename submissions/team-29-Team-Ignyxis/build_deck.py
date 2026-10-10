from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_THEME_COLOR

OUT='/mnt/data/PS01_Search_Term_Triage/DevSprint_PS01_Presentation.pptx'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
NAVY=RGBColor(17,24,39); INK=RGBColor(30,41,59); MUTED=RGBColor(100,116,139); LIGHT=RGBColor(245,247,251); WHITE=RGBColor(255,255,255); MINT=RGBColor(16,185,129); BLUE=RGBColor(59,130,246); RED=RGBColor(239,68,68); AMBER=RGBColor(245,158,11); BORDER=RGBColor(226,232,240)

def bg(slide, color=WHITE):
    sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,prs.slide_height)
    sh.fill.solid(); sh.fill.fore_color.rgb=color; sh.line.fill.background()
    slide.shapes._spTree.remove(sh._element); slide.shapes._spTree.insert(2,sh._element)

def textbox(slide,x,y,w,h,text,size=18,color=INK,bold=False,font='Aptos',align=None, valign=MSO_ANCHOR.TOP, margin=0.04):
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(margin); tf.margin_right=Inches(margin); tf.margin_top=Inches(margin); tf.margin_bottom=Inches(margin); tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.text=text
    if align is not None:p.alignment=align
    for r in p.runs:
        r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return box

def rect(slide,x,y,w,h,fill, radius=True, line=None):
    typ=MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh=slide.shapes.add_shape(typ, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line: sh.line.color.rgb=line
    else: sh.line.fill.background()
    return sh

def title(slide,kicker,heading,sub=None):
    textbox(slide,.62,.34,12,.3,kicker.upper(),10,MUTED,True)
    textbox(slide,.62,.74,12, .56,heading,28,NAVY,True)
    if sub: textbox(slide,.66,1.4,12,.42,sub,13,MUTED)

def bullets(slide,x,y,w,items,size=16,color=INK,gap=0.55):
    for i,item in enumerate(items):
        yy=y+i*gap
        rect(slide,x,yy+.06,.12,.12,MINT,True)
        textbox(slide,x+.25,yy,w-.25,gap-.04,item,size,color)

def footer(slide,num):
    textbox(slide,.65,7.12,10,.18,'DEVSPRINT 2K26  •  PS-01 SEARCH TERM TRIAGE ENGINE',8,MUTED,True)
    textbox(slide,12.15,7.09,.5,.22,str(num).zfill(2),9,MUTED,True,align=PP_ALIGN.RIGHT)

# Slide 1
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s,NAVY)
rect(s,.72,.75,.14,5.7,MINT,True)
textbox(s,1.15,1.0,10,.32,'DEVSPRINT 2K26  /  FLAGSHIP HACKATHON',11,RGBColor(167,243,208),True)
textbox(s,1.12,1.68,10.8,1.52,'Search Term\nTriage Engine',40,WHITE,True)
textbox(s,1.18,3.55,10.4,.82,'Turning 25,000 search-term rows into explainable, review-ready decisions.',20,RGBColor(203,213,225))
for i,(lab,col) in enumerate([('PROMOTE',MINT),('BLOCK',RED),('IGNORE',BLUE)]):
    rect(s,1.18+i*1.75,5.05,1.48,.49,RGBColor(31,41,55),True)
    textbox(s,1.2+i*1.75,5.16,1.44,.18,lab,10,col,True,align=PP_ALIGN.CENTER)
textbox(s,1.18,6.65,10,.22,'Kestrel Business Funding  •  Explainable decision support prototype',10,RGBColor(148,163,184))

# Slide 2
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s,'01  /  THE PROBLEM','25,000 queries. Limited human attention.','Search-term reports mix valuable business-loan intent with irrelevant or unsupported searches.')
# data cards
cards=[('25,000','search-term rows',BLUE),('17,526','unique phrases',MINT),('80.5%','rows with zero conversions',AMBER),('64.3%','rows with fewer than 5 clicks',RED)]
for i,(big,lab,col) in enumerate(cards):
    x=.68+i*3.12; rect(s,x,2.05,2.86,1.25,LIGHT,True,BORDER)
    textbox(s,x+.2,2.25,2.5,.45,big,25,col,True)
    textbox(s,x+.2,2.82,2.5,.24,lab,11,MUTED)
rect(s,.68,3.72,5.8,2.35,NAVY,True)
textbox(s,1.0,4.02,5.1,.4,'Business sells',18,WHITE,True)
bullets(s,1.0,4.58,5.0,['Working capital and business loans','Equipment finance and business lines of credit','Merchant cash advances and SBA loans'],13,RGBColor(226,232,240),.43)
rect(s,6.78,3.72,5.85,2.35,LIGHT,True,BORDER)
textbox(s,7.08,4.02,5.1,.4,'Business does NOT sell',18,NAVY,True)
bullets(s,7.08,4.58,5.0,['Personal loans, mortgages, student loans','Grants or job opportunities','Unrelated downloads and templates'],13,INK,.43)
textbox(s,.75,6.45,11.9,.3,'Core challenge: relevance + measured performance + uncertainty — not clicks alone.',14,NAVY,True)
footer(s,2)

# Slide 3
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s,'02  /  SOLUTION','An explainable triage workflow','The engine processes each row, keeps its original context, and gives a reason for every recommendation.')
steps=[('01','Validate data','Required columns, numeric fields, missing/invalid values.'),('02','Classify intent','Supported product, unsupported offering, unrelated or ambiguous.'),('03','Evaluate evidence','Clicks, cost, conversions, conversion value, ROAS and sample size.'),('04','Recommend action','PROMOTE, BLOCK or IGNORE with confidence and review flag.'),('05','Export & inspect','Searchable dashboard, campaign summaries and CSV downloads.')]
for i,(n,h,d) in enumerate(steps):
    yy=1.95+i*.9
    rect(s,.85,yy,.63,.58,NAVY,True)
    textbox(s,.88,yy+.15,.55,.18,n,12,WHITE,True,align=PP_ALIGN.CENTER)
    textbox(s,1.78,yy+.02,2.35,.3,h,16,NAVY,True)
    textbox(s,4.1,yy+.02,8.3,.48,d,14,MUTED)
    if i<4: rect(s,1.12,yy+.6,.08,.26,BORDER,False)
rect(s,.86,6.42,11.55,.42,RGBColor(236,253,245),True)
textbox(s,1.05,6.52,11.15,.18,'Decision support only: no automatic changes to live advertising accounts.',11,RGBColor(6,95,70),True)
footer(s,3)

# Slide 4
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s,'03  /  DECISION LOGIC','Three actions. Evidence-backed reasoning.','Rules are transparent and thresholds are configurable; uncertain rows are routed for human review.')
for i,(lab,col,head,desc) in enumerate([
 ('PROMOTE',MINT,'Relevant + sufficient evidence','Supported business financing intent, minimum clicks, multiple conversions and ROAS above the configured threshold.'),
 ('BLOCK',RED,'Clearly unsupported / irrelevant','Personal loans, mortgages, grants, job intent or unrelated templates; review before applying negatives.'),
 ('IGNORE',BLUE,'Uncertain or not yet proven','Thin samples, ambiguous intent, or performance that does not justify a confident action.')]):
    y=2.0+i*1.35; rect(s,.8,y,1.52,.56,LIGHT,True,BORDER); textbox(s,.87,y+.17,1.38,.18,lab,10,col,True,align=PP_ALIGN.CENTER)
    textbox(s,2.7,y+.02,3.7,.31,head,16,NAVY,True)
    textbox(s,6.45,y-.02,5.7,.68,desc,12,MUTED)
rect(s,.85,6.2,11.6,.56,RGBColor(255,251,235),True)
textbox(s,1.05,6.37,11.2,.2,'Important: a query with 2 clicks and 1 conversion is not a proven winner.',12,RGBColor(146,64,14),True)
footer(s,4)

# Slide 5
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s,'04  /  EXAMPLES','What the engine is designed to catch','Illustrative examples from the supplied synthetic dataset and advertiser profile.')
examples=[
 ('“reliable business grants for women for truckrs”','BLOCK','Grant intent is outside the lender’s offering, even though clicks/conversions may be recorded.',RED),
 ('“top working capital loan for restaurants”','PROMOTE*','Matches a supported financing product; promotion still depends on sufficient performance evidence.',MINT),
 ('“reliable SBA loan requirements calculator”','IGNORE*','Potentially relevant, but low volume can make the performance decision uncertain.',BLUE)]
for i,(term,decision,why,col) in enumerate(examples):
    y=1.95+i*1.42; rect(s,.78,y,11.8,1.12,LIGHT,True,BORDER)
    textbox(s,1.05,y+.16,7.3,.28,term,15,NAVY,True)
    rect(s,9.0,y+.15,1.32,.42,WHITE,True,BORDER); textbox(s,9.05,y+.27,1.22,.15,decision,10,col,True,align=PP_ALIGN.CENTER)
    textbox(s,1.05,y+.55,10.9,.38,why,11,MUTED)
textbox(s,.9,6.43,11.4,.34,'*Final action depends on the configured thresholds and evaluated row metrics.',10,MUTED)
footer(s,5)

# Slide 6
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s,'05  /  PROTOTYPE','Dashboard and implementation','A small, local-first application designed to be runnable during judging.')
# architectural flow
arch=[('CSV dataset','25k rows'),('Python + Pandas','validation + metrics'),('Triage engine','rules + confidence'),('Streamlit UI','filters + KPIs'),('CSV export','review queue')]
for i,(head,sub) in enumerate(arch):
    x=.65+i*2.58; rect(s,x,2.0,2.25,1.15,LIGHT,True,BORDER)
    textbox(s,x+.12,2.28,2.0,.26,head,13,NAVY,True,align=PP_ALIGN.CENTER)
    textbox(s,x+.12,2.69,2.0,.22,sub,10,MUTED,align=PP_ALIGN.CENTER)
    if i<4: textbox(s,x+2.28,2.34,.26,.3,'→',18,MINT,True,align=PP_ALIGN.CENTER)
rect(s,.75,3.75,5.75,2.25,NAVY,True)
textbox(s,1.05,4.05,5.1,.32,'Included features',17,WHITE,True)
bullets(s,1.05,4.53,5.0,['Search and filter all decisions','Campaign-level spend / value summaries','Confidence, rationale and human-review flags','Download filtered or complete CSV'],12,RGBColor(226,232,240),.33)
rect(s,6.8,3.75,5.75,2.25,LIGHT,True,BORDER)
textbox(s,7.1,4.05,5.0,.32,'Tech stack',17,NAVY,True)
bullets(s,7.1,4.53,5.0,['Python + Pandas + NumPy','Streamlit dashboard','Rule-based intent matching','No external API required by the baseline'],12,INK,.33)
footer(s,6)

# Slide 7
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s,'06  /  EVALUATION','Validate before acting on recommendations','Success means useful decisions—not merely producing a label for every row.')
cols=[('QUALITY CHECKS',MINT,['Audit unsupported-product blocks','Spot-check promoted terms','Review low-confidence rows','Inspect conversion attribution anomalies']),('METRICS TO TRACK',BLUE,['Precision of BLOCK recommendations','Promotion quality / observed ROAS','Human-review rate','Runtime and CSV export integrity']),('LIMITATIONS',AMBER,['Rule-based baseline, not trained ML','Keyword patterns may miss misspellings','Synthetic conversion data can be noisy','Thresholds need business validation'])]
for i,(head,col,items) in enumerate(cols):
    x=.75+i*4.2; rect(s,x,2.0,3.85,3.5,LIGHT,True,BORDER); textbox(s,x+.25,2.28,3.35,.28,head,12,col,True)
    bullets(s,x+.25,2.85,3.25,items,12,INK,.55)
rect(s,.85,6.05,11.65,.65,NAVY,True)
textbox(s,1.12,6.26,11.1,.22,'Next: validate against event-specific rubric, demo key cases, then freeze and package the submission.',12,WHITE,True)
footer(s,7)

prs.save(OUT)
print(OUT)
